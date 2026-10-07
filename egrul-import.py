#!/usr/bin/env python3
"""
Загрузка таблицы с ИНН в очередь ЕГРЮЛ.

Принимает .xlsx/.xlsm или .csv/.txt, сам находит колонку с ИНН (по названию,
а если названий нет — по содержимому) и складывает ИНН в egrul.csv. Уже
разобранные строки не трогает: файл можно доливать сколько угодно раз.

    python egrul-import.py --file /путь/таблица.xlsx
"""
import csv
import io
import re
import sys
from pathlib import Path

import tglib
from tglib import say

HEAD = ["inn", "status", "company", "lpr", "lpr_role", "lpr_inn",
        "phone", "address", "activity", "note", "at", "by"]

FILE = tglib.DIR / "egrul.csv"
INN_COL = re.compile(r"\bинн\b", re.I)
INN_VAL = re.compile(r"^\d{10}$|^\d{12}$")


def cells_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            yield ["" if c is None else str(c).strip() for c in row]


def cells_csv(path):
    raw = Path(path).read_bytes().decode("utf-8-sig", "replace")
    # разделитель у выгрузок бывает любой: смотрим первую строку
    head = raw.split("\n", 1)[0]
    sep = max(";,\t|", key=head.count) if any(c in head for c in ";,\t|") else ","
    for row in csv.reader(io.StringIO(raw), delimiter=sep):
        yield [c.strip() for c in row]


def innsOf(path):
    """ИНН из таблицы, по порядку и без повторов."""
    rows = list(cells_xlsx(path) if str(path).lower().endswith((".xlsx", ".xlsm"))
                else cells_csv(path))
    if not rows:
        return []

    # колонка с ИНН: сперва по заголовку, иначе — та, где больше всего значений
    # похоже на ИНН. Второй способ спасает выгрузки вообще без шапки
    col = next((i for i, c in enumerate(rows[0]) if INN_COL.search(c)), None)
    body = rows[1:] if col is not None else rows
    if col is None:
        best, hits = None, 0
        width = max(len(r) for r in rows)
        for i in range(width):
            n = sum(1 for r in rows if i < len(r) and INN_VAL.match(re.sub(r"\D", "", r[i])))
            if n > hits:
                best, hits = i, n
        if best is None:
            return []
        col, body = best, rows
        say(f"колонки «ИНН» в шапке нет — беру {col + 1}-ю: в ней {hits} похожих значений")

    out = []
    for r in body:
        if col >= len(r):
            continue
        v = re.sub(r"\D", "", r[col])
        if INN_VAL.match(v) and v not in out:
            out.append(v)
    return out


def main():
    path = tglib.arg("file", "")
    if not path or not Path(path).exists():
        raise SystemExit("нужен --file с таблицей")

    inns = innsOf(path)
    if not inns:
        say("ИНН в таблице не нашёл — проверь, та ли колонка")
        tglib.state(done=0, left=0, note="ИНН не найдены")
        return

    tglib.ensure_head(FILE, HEAD)
    have = {r.get("inn") for r in tglib.read_csv(FILE)}
    added = 0
    for inn in inns:
        if inn in have:
            continue
        tglib.append_row(FILE, [inn] + [""] * (len(HEAD) - 1), HEAD)
        added += 1

    say(f"в таблице ИНН: {len(inns)}; добавлено новых: {added}; уже были: {len(inns) - added}")
    tglib.state(done=added, left=len(inns) - added, note=f"добавлено {added}")


main()
