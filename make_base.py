#!/usr/bin/env python3
"""
Собирает рабочую базу База_Telegram.xlsx из numbers.csv,
подмешивая результаты проверки из results.csv (если он есть).

Запускать после каждого прогона check-batch.mjs — база обновится.

    ./venv/bin/python make_base.py
"""
import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DIR = Path(__file__).parent
SRC = DIR / "numbers.csv"
RES = DIR / "results.csv"
OUT = DIR / "База_Telegram.xlsx"

HEAD = ["Кто", "Звонков", "Последний звонок", "Длит. всего, с",
        "Telegram", "Имя в TG", "Проверен"]


def whole(v):
    """
    Число из ячейки. У людей, взятых из чатов, звонков и длительности нет
    вовсе — там пусто, и это не ошибка: в базе рядом с номерами теперь живут
    получатели без номера (ключ @username или id:<id>).
    """
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0

GREEN = PatternFill("solid", fgColor="C6EFCE")
GREY = PatternFill("solid", fgColor="EDEDED")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
HEAD_FILL = PatternFill("solid", fgColor="305496")


def read_csv(p: Path) -> list[dict]:
    if not p.exists():
        return []
    with p.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def main() -> None:
    base = read_csv(SRC)
    checked = {r["phone"]: r for r in read_csv(RES)}

    wb = Workbook()
    ws = wb.active
    ws.title = "База"

    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")

    stats = {"ЕСТЬ": 0, "нет": 0, "": 0}
    for r in base:
        res = checked.get(r["phone"])
        if res is None:
            status, name, when = "", "", ""
        elif res.get("tg") == "true":
            status, name, when = "ЕСТЬ", res.get("name", ""), res.get("checked_at", "")[:19]
        elif res.get("tg") == "false":
            status, name, when = "нет", "", res.get("checked_at", "")[:19]
        else:
            status, name, when = "?", "", res.get("checked_at", "")[:19]

        stats[status if status in stats else ""] += 1
        ws.append([r["phone"], whole(r.get("calls")), r.get("last_call", ""),
                   whole(r.get("total_sec")), status, name, when])

        row = ws.max_row
        if status == "ЕСТЬ":
            ws.cell(row, 5).fill = GREEN
        elif status == "нет":
            ws.cell(row, 5).fill = GREY
        elif status == "?":
            ws.cell(row, 5).fill = YELLOW

    for col, width in zip(range(1, 8), (16, 10, 20, 15, 12, 24, 20)):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    wb.save(OUT)

    from_chats = sum(1 for r in base if not str(r.get("phone", "")).startswith("+"))
    print(f"получателей в базе:  {len(base)}"
          + (f" (из них без номера, из чатов: {from_chats})" if from_chats else ""))
    print(f"  проверено:         {len(checked)}")
    print(f"    есть в Telegram: {stats['ЕСТЬ']}")
    print(f"    нет:             {stats['нет']}")
    print(f"  осталось:          {stats['']}")
    print(f"\n-> {OUT}")


if __name__ == "__main__":
    main()
